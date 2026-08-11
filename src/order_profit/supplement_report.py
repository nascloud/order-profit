"""生成 PROVISIONAL 状态的成本运费回填工作簿。"""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from order_profit.errors import OutputError
from order_profit.normalize import is_missing

if TYPE_CHECKING:
    from order_profit.inputs import LoadedTable
    from order_profit.pipeline import ValidationResult

HEADER_FILL = PatternFill("solid", fgColor="17324D")
HEADER_FONT = Font(name="Arial", color="FFFFFF", bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
INPUT_FONT = Font(name="Arial", color="0000FF")
BODY_FONT = Font(name="Arial", color="000000")

COLUMNS = [
    "主订单编号",
    "商品ID",
    "商品标题",
    "商品属性",
    "用户付款金额",
    "商家编码",
    "明细序号",
    "货品成本",
    "运费",
    "补充状态",
    "缺失项目",
    "购买数量",
    "物流单号",
    "订单状态",
    "当前核算成本",
    "当前核算运费",
    "成本来源",
    "运费来源",
]
COLUMN_INDEX = {name: index + 1 for index, name in enumerate(COLUMNS)}


def _missing_rows(result: ValidationResult) -> pd.DataFrame:
    frames = []
    for order_type, frame in (
        ("正常订单", result.analysis.profit.normal),
        ("关闭已发货", result.analysis.profit.closed_shipped),
    ):
        missing = frame[frame["核算状态"] != "已核算"].copy()
        if not missing.empty:
            missing["待补订单类型"] = order_type
            missing["待补缺失项目"] = [
                "成本+运费"
                if is_missing(row.get("核算成本_分")) and is_missing(row.get("核算运费_分"))
                else ("成本" if is_missing(row.get("核算成本_分")) else "运费")
                for _, row in missing.iterrows()
            ]
            frames.append(missing)
    missing_code = result.partitions.missing_merchant_code.copy()
    if not missing_code.empty:
        missing_code["待补订单类型"] = "无商家编码"
        missing_code["待补缺失项目"] = "成本+运费"
        for column in ("核算成本_分", "核算运费_分", "成本来源", "运费来源"):
            if column not in missing_code.columns:
                missing_code[column] = None
        frames.append(missing_code)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def _active_supplements(supplements: LoadedTable | None) -> pd.DataFrame:
    if supplements is None or supplements.data.empty:
        return pd.DataFrame()
    return supplements.data[
        supplements.data["补充状态"].isin({"", "已确认", "已同步"})
    ].copy()


def _confirmed_rows(
    active: pd.DataFrame,
    *,
    pending_detail_numbers: set[int],
) -> list[list[object]]:
    if active.empty:
        return []
    rows: list[list[object]] = []
    for _, row in active.iterrows():
        if is_missing(row.get("货品成本_分")) and is_missing(row.get("运费_分")):
            continue
        detail_number = (
            None if is_missing(row.get("明细序号")) else int(row["明细序号"])
        )
        if detail_number in pending_detail_numbers:
            continue
        rows.append(
            [
                str(row.get("主订单编号", "")),
                str(row.get("商品ID", "")),
                str(row.get("商品标题", "")),
                str(row.get("商品属性", "")),
                (
                    None
                    if is_missing(row.get("买家应付货款_分"))
                    else int(row["买家应付货款_分"]) / 100
                ),
                str(row.get("商家编码", "")),
                detail_number,
                None if is_missing(row.get("货品成本_分")) else int(row["货品成本_分"]) / 100,
                None if is_missing(row.get("运费_分")) else int(row["运费_分"]) / 100,
                "已确认",
                "已应用，留存供后续重算",
                None if is_missing(row.get("购买数量")) else int(row["购买数量"]),
                str(row.get("物流单号", "")),
                str(row.get("订单状态", "")),
                None,
                None,
                "人工补充留存",
                "人工补充留存",
            ]
        )
    return rows


def write_pending_supplements(
    result: ValidationResult,
    path: Path,
    *,
    supplements: LoadedTable | None = None,
) -> int:
    """写出可再次导入的累计人工补充表，返回仍待补行数。"""

    missing = _missing_rows(result)
    active = _active_supplements(supplements)
    pending_detail_numbers = (
        set(missing["source_row_number"].astype(int)) if not missing.empty else set()
    )
    confirmed = _confirmed_rows(
        active,
        pending_detail_numbers=pending_detail_numbers,
    )
    if missing.empty and not confirmed:
        return 0
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "人工补充表"
    worksheet.append(COLUMNS)
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    last_column = get_column_letter(len(COLUMNS))
    worksheet.auto_filter.ref = (
        f"A1:{last_column}{len(confirmed) + len(missing) + 1}"
    )

    for values in confirmed:
        worksheet.append(values)

    first_pending_row = len(confirmed) + 2
    for output_row, (_, row) in enumerate(missing.iterrows(), start=first_pending_row):
        cost_missing = is_missing(row.get("核算成本_分"))
        freight_missing = is_missing(row.get("核算运费_分"))
        missing_items = str(row["待补缺失项目"])
        values = [
            str(row.get("主订单编号", "")),
            str(row.get("商品ID", "")),
            str(row.get("商品标题", "")),
            str(row.get("商品属性", "")),
            (
                None
                if is_missing(row.get("买家应付货款_分"))
                else int(row["买家应付货款_分"]) / 100
            ),
            str(row.get("商家编码", "")),
            int(row["source_row_number"]),
            None,
            None,
            "草稿",
            missing_items,
            int(row.get("购买数量_标准", 0)),
            str(row.get("物流单号", "")),
            str(row.get("订单状态", "")),
            None if cost_missing else int(row["核算成本_分"]) / 100,
            None if freight_missing else int(row["核算运费_分"]) / 100,
            str(row.get("成本来源", "")),
            str(row.get("运费来源", "")),
        ]
        worksheet.append(values)
    input_columns = [
        COLUMN_INDEX["货品成本"],
        COLUMN_INDEX["运费"],
        COLUMN_INDEX["补充状态"],
    ]
    text_columns = [
        COLUMN_INDEX["主订单编号"],
        COLUMN_INDEX["商品ID"],
        COLUMN_INDEX["商家编码"],
        COLUMN_INDEX["明细序号"],
        COLUMN_INDEX["物流单号"],
    ]
    money_columns = [
        COLUMN_INDEX["用户付款金额"],
        COLUMN_INDEX["货品成本"],
        COLUMN_INDEX["运费"],
        COLUMN_INDEX["当前核算成本"],
        COLUMN_INDEX["当前核算运费"],
    ]
    for output_row in range(2, worksheet.max_row + 1):
        for column in range(1, len(COLUMNS) + 1):
            worksheet.cell(output_row, column).font = BODY_FONT
        for column in input_columns:
            cell = worksheet.cell(output_row, column)
            cell.fill = INPUT_FILL
            cell.font = INPUT_FONT
        for column in text_columns:
            worksheet.cell(output_row, column).number_format = "@"
        for column in money_columns:
            worksheet.cell(output_row, column).number_format = '#,##0.00;[Red]-#,##0.00'

    status_validation = DataValidation(
        type="list",
        formula1='"草稿,已确认,作废"',
        allow_blank=False,
    )
    status_validation.error = "请选择草稿、已确认或作废"
    status_validation.errorTitle = "补充状态无效"
    worksheet.add_data_validation(status_validation)
    status_column = get_column_letter(COLUMN_INDEX["补充状态"])
    status_validation.add(
        f"{status_column}2:{status_column}{len(confirmed) + len(missing) + 1}"
    )
    comments = {
        "货品成本": "黄色单元格：填写该订单明细的总货品成本，单位元。",
        "运费": "黄色单元格：填写该订单明细的总运费，单位元。",
        "补充状态": "填写完成后将草稿改为已确认；草稿行不会参与核算。",
        "用户付款金额": "订单表中的买家应付货款，仅供核对，不可编辑。",
    }
    for column_name, text in comments.items():
        column_letter = get_column_letter(COLUMN_INDEX[column_name])
        worksheet[f"{column_letter}1"].comment = Comment(text, "order-profit")
    widths = {
        "A": 24,
        "B": 18,
        "C": 46,
        "D": 46,
        "E": 16,
        "F": 28,
        "G": 12,
        "H": 14,
        "I": 14,
        "J": 12,
        "K": 14,
        "L": 12,
        "M": 20,
        "N": 24,
        "O": 16,
        "P": 16,
        "Q": 24,
        "R": 24,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    worksheet.sheet_view.showGridLines = False

    guide = workbook.create_sheet("填写说明")
    guide_rows = [
        ["步骤", "操作"],
        ["1", "在人工补充表中筛选缺失项目。"],
        ["2", "按缺失项目填写黄色成本和/或运费；原商家编码为空的明细必须同时填写成本和运费。"],
        ["3", "金额为当前订单明细总金额；货品成本和运费均以元为单位，不再乘购买数量。"],
        ["4", "填写完成后把补充状态从草稿改为已确认。"],
        ["5", "保存文件；不要修改主订单编号、商品ID、原商家编码和明细序号。"],
        ["6", "回到桌面端，在“人工补充表”中选择本文件，重新生成经营复盘。"],
        ["说明", "草稿和作废行不参与核算；已确认行必须至少填写成本或运费之一，无商家编码行必须两项均填写。"],
        ["说明", "明细序号是原订单 Excel 行号，用于精确定位重复基础键。"],
    ]
    for row in guide_rows:
        guide.append(row)
    for cell in guide[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    guide.column_dimensions["A"].width = 14
    guide.column_dimensions["B"].width = 90
    guide.sheet_view.showGridLines = False

    workbook.properties.creator = "order-profit"
    workbook.properties.title = "待补成本运费"
    workbook.save(path)
    validate_pending_supplements(path, len(confirmed) + len(missing))
    return len(missing)


def validate_pending_supplements(path: Path, expected_rows: int) -> None:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise OutputError(f"待补成本运费表写后读回失败：{exc}", code="SUPPLEMENT_REPORT_READBACK_FAILED") from exc
    try:
        if workbook.sheetnames[:2] != ["人工补充表", "填写说明"]:
            raise OutputError("待补成本运费表工作表结构错误", code="SUPPLEMENT_REPORT_SHEETS_INVALID")
        worksheet = workbook["人工补充表"]
        headers = [cell.value for cell in worksheet[1]]
        if headers != COLUMNS or worksheet.max_row - 1 != expected_rows:
            raise OutputError("待补成本运费表字段或行数不一致", code="SUPPLEMENT_REPORT_CONTENT_INVALID")
    finally:
        workbook.close()
