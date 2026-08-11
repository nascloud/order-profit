"""Excel 工作簿生成、格式设置与写后校验。"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING, Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from order_profit.config import RunConfig
from order_profit.errors import OutputError
from order_profit.normalize import is_missing

if TYPE_CHECKING:
    from order_profit.pipeline import ValidationResult

ERROR_TOKENS = ("#VALUE!", "#DIV/0!", "#NAME?", "#REF!", "#N/A")
HEADER_FILL = PatternFill("solid", fgColor="17324D")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBHEADER_FILL = PatternFill("solid", fgColor="EAF0F5")
MONEY_FORMAT = '#,##0.00;[Red]-#,##0.00'
PERCENT_FORMAT = '0.00%'
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
INPUT_FONT = Font(name="Arial", color="0000FF")
BODY_FONT = Font(name="Arial", color="000000")


def _safe_excel_value(value: Any) -> Any:
    if is_missing(value):
        return None
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date, pd.Timestamp)):
        parsed = value.to_pydatetime() if isinstance(value, pd.Timestamp) else value
        if isinstance(parsed, datetime) and parsed.tzinfo is not None:
            return parsed.isoformat()
        return parsed
    if isinstance(value, (tuple, list, set)):
        return "、".join(map(str, value))
    if isinstance(value, dict):
        return str(value)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _display_frame(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    rename: dict[str, str] = {}
    for column in result.columns:
        if column.endswith("_分"):
            result[column] = result[column].map(
                lambda value: None if is_missing(value) else int(value) / 100
            )
            rename[column] = column.removesuffix("_分")
    return result.rename(columns=rename)


def _write_frame(workbook: Workbook, name: str, frame: pd.DataFrame) -> None:
    worksheet = workbook.create_sheet(name[:31])
    display = _display_frame(frame)
    columns = [str(column) for column in display.columns]
    if not columns:
        worksheet.append(["说明"])
        worksheet.append(["本次运行无数据"])
    else:
        worksheet.append(columns)
        for row in display.itertuples(index=False, name=None):
            safe_row = [_safe_excel_value(value) for value in row]
            for value in safe_row:
                if isinstance(value, str) and any(token in value for token in ERROR_TOKENS):
                    raise OutputError(
                        f"工作表 {name} 的业务文本包含 Excel 错误标记：{value}",
                        code="EXCEL_ERROR_TOKEN",
                    )
            worksheet.append(safe_row)
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, column in enumerate(worksheet.iter_cols(1, worksheet.max_column), start=1):
        header = str(column[0].value or "")
        width = max(len(header) * 1.7 + 3, 10)
        for cell in list(column)[1:101]:
            if cell.value is not None:
                width = max(width, min(len(str(cell.value)) * 1.3 + 2, 44))
            if header.endswith(("金额", "成本", "运费", "利润", "损失", "扣点", "税赋", "花费", "销售额", "退款")):
                cell.number_format = MONEY_FORMAT
            elif header.endswith(("率", "比例")) or header == "推广投产比":
                cell.number_format = PERCENT_FORMAT
            if any(token in header for token in ("ID", "编号", "单号", "编码")):
                cell.number_format = "@"
        worksheet.column_dimensions[get_column_letter(index)].width = min(width, 46)
    worksheet.sheet_view.showGridLines = False


def _summary_frame(result: ValidationResult) -> pd.DataFrame:
    rows = []
    for metric, value in result.analysis.summary.items():
        unit = ""
        display_value: Any = value
        if metric.endswith("_分"):
            metric = metric.removesuffix("_分")
            display_value = int(value) / 100
            unit = "元"
        elif metric.endswith("率"):
            display_value = None if is_missing(value) else float(value) * 100
            unit = "%"
        rows.append({"指标": metric, "值": display_value, "单位": unit})
    return pd.DataFrame(rows)


def _recommendations(result: ValidationResult) -> pd.DataFrame:
    links = result.analysis.links
    if links.empty:
        return pd.DataFrame([{"优先级": "低", "行动": "本期无有效商品链接", "对象": "店铺"}])
    rows: list[dict[str, object]] = []
    for _, row in links.iterrows():
        if row["运营动作"] == "保持观察":
            continue
        rows.append(
            {
                "优先级": row["运营动作优先级"],
                "行动": row["运营动作"],
                "对象": row["商品ID"],
                "商品标题": row["商品标题"],
                "经营贡献利润": int(row["含推广经营贡献利润_分"]) / 100,
            }
        )
    return pd.DataFrame(rows or [{"优先级": "低", "行动": "保持观察", "对象": "店铺"}])


def _rules_frame(result: ValidationResult, config: RunConfig) -> pd.DataFrame:
    config_rows = [{"项目": key, "值": _safe_excel_value(value)} for key, value in config.to_dict().items()]
    config_rows.extend(
        [
            {"项目": "运行状态", "值": result.reconciliation.status.value},
            {"项目": "费用计费基数", "值": "MAX(实收金额-有效退款, 0)"},
            {"项目": "金额口径", "值": "逐行 ROUND_HALF_UP 到分，整数分汇总"},
            {"项目": "推广归因成交", "值": "仅用于投产分析，不计入订单销售额"},
            {"项目": "产品推广数据", "值": "链接级推广按产品已付款销售额确定性分摊估算"},
            {"项目": "平台扣点与税赋", "值": "管理估算值，非平台实际结算费用"},
            {"项目": "SABC阈值", "值": f"S={config.s_profit_share}，A={config.a_profit_share}"},
            {"项目": "Excel事实来源", "值": "Python PipelineResult；工作簿写入已计算值"},
        ]
    )
    return pd.DataFrame(config_rows)


def write_excel(result: ValidationResult, config: RunConfig, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    partitions = result.partitions
    analysis = result.analysis
    sheets = [
        ("01_经营总览", _summary_frame(result)),
        ("02_每日经营", analysis.daily),
        ("03_产品分析", analysis.products),
        ("04_链接分析", analysis.links),
        ("正常订单明细", analysis.profit.normal),
        ("关闭已发货", analysis.profit.closed_shipped),
        ("关闭无物流已付款", analysis.profit.closed_unshipped),
        ("未付款关闭审计", partitions.unpaid_closed),
        ("付款状态异常", partitions.payment_anomaly),
        ("无商家编码", partitions.missing_merchant_code),
        ("成本运费补充记录", result.costs.supplement_audit),
        ("清洗删除明细", partitions.deleted),
        ("超出日期范围", partitions.out_of_period),
        ("其他状态", partitions.other_status),
        ("SKU分析", analysis.sku),
        ("售后退款复核", analysis.profit.after_sales_reconciliation),
        ("推广明细", analysis.promotion.detail),
        ("推广每日汇总", analysis.promotion.daily),
        ("推广商品分析", analysis.promotion.products),
        ("关键因素与建议", _recommendations(result)),
        ("规则说明", _rules_frame(result, config)),
    ]
    for name, frame in sheets:
        _write_frame(workbook, name, frame)
    workbook.properties.creator = "order-profit"
    workbook.properties.title = "订单成本利润分析"
    workbook.save(path)
    validate_excel_report(path, result)


def validate_excel_report(path: Path, result: ValidationResult) -> None:
    try:
        workbook = load_workbook(path, data_only=False, read_only=True)
    except Exception as exc:
        raise OutputError(f"Excel 写后读回失败：{exc}", code="EXCEL_READBACK_FAILED") from exc
    try:
        required = {"01_经营总览", "02_每日经营", "03_产品分析", "04_链接分析", "正常订单明细", "规则说明"}
        missing = required.difference(workbook.sheetnames)
        if missing:
            raise OutputError(
                f"Excel 缺少工作表：{'、'.join(sorted(missing))}",
                code="EXCEL_SHEET_MISSING",
            )
        summary_sheet = workbook["01_经营总览"]
        values = {
            str(summary_sheet.cell(row=row, column=1).value): summary_sheet.cell(row=row, column=2).value
            for row in range(2, summary_sheet.max_row + 1)
        }
        expected = int(result.analysis.summary["含推广经营贡献利润_分"]) / 100
        actual = values.get("含推广经营贡献利润")
        if actual is None or abs(float(actual) - expected) > 0.01:
            raise OutputError(
                f"Excel 汇总读回不一致：含推广经营贡献利润 expected={expected}, actual={actual}",
                code="EXCEL_SUMMARY_MISMATCH",
            )
    finally:
        workbook.close()
