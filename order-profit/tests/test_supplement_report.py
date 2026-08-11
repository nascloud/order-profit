from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import load_workbook

from order_profit.config import Platform, RunConfig
from order_profit.costs import apply_supplements
from order_profit.inputs import InputFileMetadata, LoadedTable, load_supplements
from order_profit.models import RunStatus
from order_profit.normalize import add_source_audit_columns
from order_profit.supplement_report import COLUMNS, write_pending_supplements
from order_profit.validation import normalize_supplements

TZ = ZoneInfo("Asia/Shanghai")


def _config(tmp_path: Path) -> RunConfig:
    return RunConfig(
        orders=tmp_path / "orders.xlsx",
        start_date=date(2026, 8, 1),
        end_date=date(2026, 8, 7),
        platform=Platform.TMALL,
        output_dir=tmp_path / "output",
    )


def _incomplete_result() -> SimpleNamespace:
    normal = pd.DataFrame(
        [
            {
                "source_row_number": 17,
                "主订单编号": "O1",
                "商品ID": "P1",
                "商家编码": "SKU1",
                "商品标题": "测试商品",
                "商品属性": "颜色分类:白色；规格:加厚款",
                "购买数量_标准": 2,
                "物流单号": "W1",
                "订单状态": "交易成功",
                "订单付款时间_标准": datetime(2026, 8, 1, tzinfo=TZ),
                "核算成本_分": None,
                "核算运费_分": 800,
                "成本来源": "缺失",
                "运费来源": "运费表实际运费",
                "核算状态": "成本或运费缺失",
            }
        ]
    )
    return SimpleNamespace(
        analysis=SimpleNamespace(
            profit=SimpleNamespace(normal=normal, closed_shipped=normal.iloc[0:0].copy())
        ),
        reconciliation=SimpleNamespace(status=RunStatus.PROVISIONAL),
        partitions=SimpleNamespace(missing_merchant_code=normal.iloc[0:0].copy()),
    )


def test_pending_workbook_is_directly_loadable_as_draft(tmp_path: Path) -> None:
    path = tmp_path / "待补成本运费.xlsx"
    count = write_pending_supplements(_incomplete_result(), path)

    assert count == 1
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook["人工补充表"]
    assert [cell.value for cell in worksheet[1]] == COLUMNS
    assert worksheet["E2"].value == 17
    assert worksheet["F2"].value is None
    assert worksheet["G2"].value is None
    assert worksheet["H2"].value == "草稿"
    assert worksheet["K2"].value == "颜色分类:白色；规格:加厚款"
    for coordinate in ("D2", "F2", "G2", "H2"):
        assert worksheet[coordinate].fill.fgColor.rgb.endswith("FFF2CC")
    workbook.close()

    loaded = load_supplements(path, "人工补充表")
    normalized = normalize_supplements(loaded, _config(tmp_path))
    assert normalized.data.iloc[0]["补充状态"] == "草稿"
    assert normalized.data.iloc[0]["明细序号"] == 17


def test_pending_workbook_keeps_confirmed_rows_for_next_rerun(tmp_path: Path) -> None:
    supplements = pd.DataFrame(
        [
            {
                "主订单编号": "OLD",
                "商品ID": "P0",
                "商家编码": "",
                "补充商家编码": "",
                "明细序号": 16,
                "货品成本_分": 1234,
                "运费_分": 800,
                "补充状态": "已确认",
                "商品标题": "已完成商品",
                "商品属性": "颜色:白色",
                "购买数量": 1,
                "物流单号": "W0",
                "订单状态": "交易成功",
            }
        ]
    )
    loaded = LoadedTable(
        supplements,
        InputFileMetadata(
            role="人工补充表",
            path=tmp_path / "old.xlsx",
            sha256="b" * 64,
            size_bytes=1,
            rows=1,
            sheet="人工补充表",
        ),
    )
    path = tmp_path / "待补成本运费.xlsx"

    count = write_pending_supplements(
        _incomplete_result(),
        path,
        supplements=loaded,
    )

    assert count == 1
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook["人工补充表"]
    assert worksheet.max_row == 3
    assert worksheet["A2"].value == "OLD"
    assert worksheet["F2"].value == 12.34
    assert worksheet["G2"].value == 8
    assert worksheet["H2"].value == "已确认"
    assert worksheet["I2"].value == "已应用，留存供后续重算"
    assert worksheet["A3"].value == "O1"
    assert worksheet["H3"].value == "草稿"
    workbook.close()


def test_missing_merchant_code_can_be_confirmed_by_detail_number(tmp_path: Path) -> None:
    result = _incomplete_result()
    missing_code = result.analysis.profit.normal.copy()
    missing_code["商家编码"] = ""
    missing_code["核算状态"] = None
    result.partitions = SimpleNamespace(missing_merchant_code=missing_code)
    # This synthetic result includes only the missing-code partition.
    result.analysis.profit.normal = result.analysis.profit.normal.iloc[0:0].copy()
    path = tmp_path / "待补成本运费.xlsx"
    write_pending_supplements(result, path)
    workbook = load_workbook(path)
    worksheet = workbook["人工补充表"]
    assert worksheet["I2"].value == "商家编码或成本+运费"
    worksheet["D2"] = "NEW-SKU"
    worksheet["H2"] = "已确认"
    workbook.save(path)
    workbook.close()

    loaded = load_supplements(path, "人工补充表")
    audited = add_source_audit_columns(
        loaded.data,
        path=loaded.metadata.path,
        sheet=loaded.metadata.sheet,
        file_sha256=loaded.metadata.sha256,
    )
    supplement = normalize_supplements(
        LoadedTable(audited, loaded.metadata, loaded.issues),
        _config(tmp_path),
    )
    orders_data = pd.DataFrame(
        [
            {
                "source_row_number": 17,
                "主订单编号": "O1",
                "商品ID": "P1",
                "商家编码": "",
                "货品成本_分": None,
                "运费_分": None,
            }
        ]
    )
    orders = LoadedTable(
        orders_data,
        InputFileMetadata(
            role="订单表",
            path=tmp_path / "orders.xlsx",
            sha256="a" * 64,
            size_bytes=1,
            rows=1,
            sheet="export",
        ),
    )

    applied = apply_supplements(orders, supplement, overwrite_existing=False)

    assert applied.orders.data.iloc[0]["商家编码"] == "NEW-SKU"
    assert bool(applied.audit.iloc[0]["编码已应用"])


def test_missing_merchant_code_can_be_confirmed_without_supplemental_code(
    tmp_path: Path,
) -> None:
    result = _incomplete_result()
    missing_code = result.analysis.profit.normal.copy()
    missing_code["商家编码"] = ""
    missing_code["核算状态"] = None
    result.partitions = SimpleNamespace(missing_merchant_code=missing_code)
    result.analysis.profit.normal = result.analysis.profit.normal.iloc[0:0].copy()
    path = tmp_path / "待补成本运费.xlsx"
    write_pending_supplements(result, path)
    workbook = load_workbook(path)
    worksheet = workbook["人工补充表"]
    worksheet["F2"] = 12.34
    worksheet["G2"] = 8
    worksheet["H2"] = "已确认"
    workbook.save(path)
    workbook.close()

    loaded = load_supplements(path, "人工补充表")
    audited = add_source_audit_columns(
        loaded.data,
        path=loaded.metadata.path,
        sheet=loaded.metadata.sheet,
        file_sha256=loaded.metadata.sha256,
    )
    supplement = normalize_supplements(
        LoadedTable(audited, loaded.metadata, loaded.issues),
        _config(tmp_path),
    )

    assert supplement.data.iloc[0]["补充商家编码"] == ""
    assert supplement.data.iloc[0]["货品成本_分"] == 1234
    assert supplement.data.iloc[0]["运费_分"] == 800


def test_confirmed_generated_row_applies_by_detail_number(tmp_path: Path) -> None:
    path = tmp_path / "待补成本运费.xlsx"
    write_pending_supplements(_incomplete_result(), path)
    workbook = load_workbook(path)
    worksheet = workbook["人工补充表"]
    worksheet["F2"] = 12.34
    worksheet["H2"] = "已确认"
    workbook.save(path)
    workbook.close()

    loaded = load_supplements(path, "人工补充表")
    audited_data = add_source_audit_columns(
        loaded.data,
        path=loaded.metadata.path,
        sheet=loaded.metadata.sheet,
        file_sha256=loaded.metadata.sha256,
    )
    loaded = LoadedTable(audited_data, loaded.metadata, loaded.issues)
    supplement = normalize_supplements(loaded, _config(tmp_path))
    orders_data = pd.DataFrame(
        [
            {
                "source_row_number": 17,
                "主订单编号": "O1",
                "商品ID": "P1",
                "商家编码": "SKU1",
                "货品成本_分": None,
                "运费_分": 800,
            },
            {
                "source_row_number": 18,
                "主订单编号": "O1",
                "商品ID": "P1",
                "商家编码": "SKU1",
                "货品成本_分": None,
                "运费_分": 900,
            },
        ]
    )
    orders = LoadedTable(
        data=orders_data,
        metadata=InputFileMetadata(
            role="订单表",
            path=tmp_path / "orders.xlsx",
            sha256="a" * 64,
            size_bytes=1,
            rows=2,
            sheet="export",
        ),
    )

    result = apply_supplements(orders, supplement, overwrite_existing=False)

    assert result.orders.data.iloc[0]["补充货品成本_分"] == 1234
    assert result.orders.data.iloc[1]["补充货品成本_分"] is None
    assert result.audit.iloc[0]["明细序号"] == 17
    assert result.audit.iloc[0]["匹配订单源行"] == 17
